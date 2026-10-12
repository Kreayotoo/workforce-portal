#!/usr/bin/env python3
"""Build digital-twin HTML for each *_All_Utilities_Quantities.xlsx block (no above-ground city)."""
import openpyxl, glob, os, re, json
import numpy as np
from scipy.interpolate import griddata
from pyproj import Transformer

UP='/mnt/user-data/uploads'
PKG='/home/claude/zone5pkg/Amaravati_Twins'
OUT=PKG+'/blocks'; os.makedirs(OUT,exist_ok=True)
TPL=open('/home/claude/dashboard_template.html').read()
LOGO=open('/home/claude/logo_b64.txt').read().strip()
T=Transformer.from_crs(32644,4326,always_xy=True)
ORDER=['water','sewer','strom','reuse','power']
LABEL={'water':'Water','sewer':'Sewer','strom':'Storm','reuse':'Reuse','power':'Power'}

def _pill(zone, blk):
    m=re.match(r'^B(\d+)$', str(blk))
    bn = m.group(1) if m else str(blk).replace('_','-')
    return f'ZONE {zone} · BLOCK {bn}'

def _manifest_entry(fn, code, zone, blk):
    h=open(fn).read()
    net=re.search(r'const NET\s*=\s*(\{.*?\});', h, re.S).group(1)
    o=json.loads(net); st=o.get('stats',{})
    pipes=sum(st[k]['count'] for k in ORDER if k in st)
    km=round(sum(st[k]['length_m'] for k in ORDER if k in st)/1000,2)
    ctr=o.get('map',{}).get('center',[0,0])
    return {'code':code,'zone':zone,'block':blk,'file':'blocks/'+os.path.basename(fn),
            'pipes':pipes,'manholes':len(o.get('nodes',[])),'length_km':km,'lat':ctr[0],'lon':ctr[1],
            'utils':{k:{'pipes':st[k]['count'],'km':round(st[k]['length_m']/1000,2)} for k in ORDER if k in st}}

def f(v):
    if v is None: return None
    s=str(v).strip().replace(',','')
    if s in ('','<None>'): return None
    m=re.match(r'^(-?\d+\.?\d*)',s)
    return float(m.group(1)) if m else None
def mm(v):  # → millimetres
    if v is None: return None
    s=str(v).strip().replace(',','')
    if s in ('','<None>'): return None
    if s.endswith('mm'): return f(s[:-2])
    if s.endswith('in'): return f(s[:-2])*25.4
    return f(s)
def meters(v):
    x=mm(v); return None if x is None else x/1000.0
def ang(v):
    if not v: return 0.0
    m=re.match(r'^(-?\d+)\D+(\d+)\D+([\d.]+)',str(v))
    if m: return float(m.group(1))+float(m.group(2))/60+float(m.group(3))/3600
    return f(v) or 0.0

def util_of(name):
    n=name.strip().lower()
    if 'reuse' in n: return 'reuse'
    if 'water' in n: return 'water'
    if 'power' in n or 'pmh' in n or 'config' in n: return 'power'
    if 'fsmh' in n or 'sewer' in n: return 'sewer'
    if 'storm' in n or 'drain' in n or 'smh' in n: return 'strom'
    return None

def classify(wb):
    pipes={}; mh={}
    for sh in wb.sheetnames:
        if sh.strip().lower()=='quantities': continue
        u=util_of(sh)
        if u is None: continue
        ws=wb[sh]; ftype=None
        for r in ws.iter_rows(min_row=1, max_row=8, values_only=True):
            cells=[str(c).strip().lower() if c is not None else '' for c in (r or [])]
            if not cells or not any(cells): continue
            if cells[0].startswith('structure name'): ftype='mh'; break
            if cells[0].startswith('pipe name'): ftype='pipe'; break
            if any('easting' in c for c in cells):
                if any(('rim' in c) or ('sump' in c) for c in cells): ftype='mh'
                elif any('invert' in c for c in cells): ftype='pipe'
                break
        if ftype=='mh': mh[u]=sh
        elif ftype=='pipe': pipes[u]=sh
    return pipes, mh

def rows_of(ws):
    allrows=list(ws.iter_rows(values_only=True))
    hidx=0
    for i,r in enumerate(allrows[:8]):
        cells=[str(c).strip().lower() if c is not None else '' for c in (r or [])]
        if not cells: continue
        if cells[0].startswith('pipe name') or cells[0].startswith('structure name'): hidx=i; break
        if any('easting' in c for c in cells) and any(('invert' in c) or ('rim' in c) or ('sump' in c) for c in cells): hidx=i; break
    hdr=[str(c).strip() if c is not None else '' for c in allrows[hidx]]
    ix={h:i for i,h in enumerate(hdr)}
    return allrows[hidx+1:], ix

def get(r,ix,name):
    i=ix.get(name); return r[i] if (i is not None and i<len(r)) else None

def build_block(path):
    base=os.path.basename(path)
    raw=re.sub(r'^[0-9]+_','',base)
    raw=re.sub(r'_All_Utilities_Quantities','',raw, flags=re.I)
    raw=re.sub(r'\.xlsx$','',raw, flags=re.I)            # e.g. 5A_B2, 5D_B1_B2, Z5D_B3, Z10_B3, Z12_B4
    code=re.sub(r'^[Zz](?=\d)','',raw)                   # strip a leading Z before a digit -> 10_B3, 5D_B3
    parts=code.split('_',1)
    zone=parts[0]; blk=parts[1] if len(parts)>1 else 'B1'
    fn=os.path.join(OUT,f'AmaravatiTwin_{code}.html')
    if os.path.exists(fn) and os.path.getsize(fn)>50000:
        print(f'  {code}: already built, skipping'); 
        return _manifest_entry(fn, code, zone, blk)
    wb=openpyxl.load_workbook(path, read_only=True, data_only=True)
    pipesh, mhsh = classify(wb)

    # ---- parse pipes ----
    nets={k:[] for k in ORDER}; dims={k:[] for k in ORDER}
    allE=[]; allN=[]
    for u,sh in pipesh.items():
        rows,ix=rows_of(wb[sh])
        for r in rows:
            if not r or r[0] in (None,''): continue
            se=f(get(r,ix,'Start Easting')); sn=f(get(r,ix,'Start Northing'))
            ee=f(get(r,ix,'End Easting'));   en=f(get(r,ix,'End Northing'))
            if None in (se,sn,ee,en): continue
            si=f(get(r,ix,'Start Invert')) or 0.0; ei=f(get(r,ix,'End Invert')) or 0.0
            iw=mm(get(r,ix,'Inner Width')) or 200.0; ih=mm(get(r,ix,'Inner Height')) or iw
            nets[u].append([se,sn,si,ee,en,ei])
            box=1 if u in ('strom','power') else 0
            dims[u].append([box, round(iw,1), round(ih,1)])
            allE+=[se,ee]; allN+=[sn,en]
    if not allE: 
        print('  ! no pipe geometry in',code); return None
    allE=np.array(allE); allN=np.array(allN)
    cE=float(allE.mean()); cN=float(allN.mean())

    # ---- relativise + stats + geo (lat/lon) ----
    networks={}; geo={}; stats={}
    for k in ORDER:
        segs=nets[k]
        networks[k]=[[round(s[0]-cE,2),round(s[1]-cN,2),round(s[2],3),
                      round(s[3]-cE,2),round(s[4]-cN,2),round(s[5],3)] for s in segs]
        if segs:
            arr=np.array(segs)
            sl,so=T.transform(arr[:,0],arr[:,1]); el,eo=T.transform(arr[:,3],arr[:,4])
            geo[k]=[[round(so[i],6),round(sl[i],6),round(eo[i],6),round(el[i],6)] for i in range(len(segs))]
            L=np.hypot(arr[:,3]-arr[:,0],arr[:,4]-arr[:,1]).sum()
            d=np.array([x[1] for x in dims[k]])
            invs=np.concatenate([arr[:,2],arr[:,5]])
            stats[k]={'label':LABEL[k],'count':len(segs),'length_m':round(float(L),1),
                'avg_dia':int(d.mean()),'max_dia':int(d.max()),'min_dia':int(d.min()),
                'min_inv':round(float(invs.min()),2),'max_inv':round(float(invs.max()),2)}
        else:
            geo[k]=[]; stats[k]={'label':LABEL[k],'count':0,'length_m':0,'avg_dia':0,'max_dia':0,'min_dia':0,'min_inv':0,'max_inv':0}

    # ---- parse structures (manholes) ----
    nodes=[]; rimpts=[]
    for u,sh in mhsh.items():
        rows,ix=rows_of(wb[sh])
        for i,r in enumerate(rows):
            if not r or r[0] in (None,''): continue
            E=f(get(r,ix,'Easting')); N=f(get(r,ix,'Northing'))
            if E is None or N is None: continue
            rim=f(get(r,ix,'Rim Elevation')); sump=f(get(r,ix,'Sump Elevation'))
            r2s=f(get(r,ix,'Rim to Sump Height'))
            if rim is None and sump is not None: rim=sump+2.0
            if rim is None: rim=f(get(r,ix,'Surface Elevation'))
            if rim is None: continue
            if sump is None: sump=rim-(r2s if r2s else 2.0)
            if sump>=rim: sump=rim-0.6
            iw=meters(get(r,ix,'Inner Width')); il=meters(get(r,ix,'Inner Length'))
            shp=str(get(r,ix,'Shape') or '').strip().lower()
            circ = shp.startswith('cyl') if shp else (u=='sewer')
            lon,lat=T.transform(E,N)
            pipes=str(get(r,ix,'Connected Pipes') or '').strip()
            nodes.append({'u':u,'sid':str(get(r,ix,'Structure Name') or (u+'-'+str(i+1))),
                'desc':str(get(r,ix,'Description') or ''),
                'x':round(E-cE,2),'z':round(N-cN,2),'E':round(E,3),'N':round(N,3),
                'inv':round(sump,2),'rim':round(rim,2),'frl':round(rim,2),
                'dia':round(iw or (1.2 if u=='strom' else 2.4),3),'len':round(il,3) if il else 0,
                'coneH':round(meters(get(r,ix,'Cone Height')) or 0,3),
                'neck':round(meters(get(r,ix,'Frame Length')) or 0.45,3),
                'rot':round(ang(get(r,ix,'Rotation Angle')),2),
                'shape':'cyl' if circ else 'box','pipes':pipes,
                'deg':max(1,len([p for p in pipes.split(',') if p.strip() and p.strip()!='<None>'])),
                'mh':True,'real':1,'lat':round(lat,6),'lon':round(lon,6)})
            rimpts.append((E,N,rim))

    # ---- road surface from manhole rim levels (no LandXML for these blocks) ----
    relx=allE-cE; relz=allN-cN
    ex0,ex1=np.percentile(relx,2),np.percentile(relx,98)
    ez0,ez1=np.percentile(relz,2),np.percentile(relz,98)
    M=50.0; x0,x1,z0,z1=ex0-M,ex1+M,ez0-M,ez1+M
    DX=5.0
    nx=min(int(np.ceil((x1-x0)/DX))+1,260); nz=min(int(np.ceil((z1-z0)/DX))+1,260)
    road=None; frl_level=None
    if len(rimpts)>=6:
        P=np.array(rimpts); P[:,0]-=cE; P[:,1]-=cN
        gx=x0+np.arange(nx)*DX; gz=z0+np.arange(nz)*DX
        GX,GZ=np.meshgrid(gx,gz)
        H=griddata((P[:,0],P[:,1]),P[:,2],(GX,GZ),method='linear')
        nanm=~np.isfinite(H)
        if nanm.any():
            Hn=griddata((P[:,0],P[:,1]),P[:,2],(GX,GZ),method='nearest'); H[nanm]=Hn[nanm]
        H=np.round(H,2)
        road={'x0':round(x0,2),'z0':round(z0,2),'dx':DX,'nx':nx,'nz':nz,
              'emin':float(H.min()),'emax':float(H.max()),'h':[float(v) for v in H.flatten()]}
        frl_level=float(round(float(np.median(H)),2))
    if frl_level is None:
        frl_level=round(float(np.mean([n['rim'] for n in nodes])) if nodes else 25.0,2)

    # ---- map factors ----
    lon0,lat0=T.transform(cE,cN); lonE,latE=T.transform(cE+1,cN); lonN,latN=T.transform(cE,cN+1)
    latlonsW,latlonsS=T.transform(cE+ex0,cN+ez0); 
    sw_lon,sw_lat=T.transform(cE+relx.min(),cN+relz.min()); ne_lon,ne_lat=T.transform(cE+relx.max(),cN+relz.max())
    NET={'meta':{'zone':zone,'block':blk,'grid':'44N','center_E':round(cE,2),'center_N':round(cN,2),'aboveground':False},
         'stats':stats,'networks':networks,'geo':geo,'dims':dims,'nodes':nodes,
         'nodes_src':'civil3d quantities workbook (pipes+manholes)','nodes_frl':True,
         'frl':{'level':frl_level,'note':'derived from manhole rim levels (no LandXML)','src':'rim-interp'},
         'map':{'center':[round(lat0,6),round(lon0,6)],
                'bounds':[[round(sw_lat,6),round(sw_lon,6)],[round(ne_lat,6),round(ne_lon,6)]],
                'lonPerM_x':lonE-lon0,'latPerM_x':latE-lat0,'lonPerM_n':lonN-lon0,'latPerM_n':latN-lat0}}
    if road: NET['road']=road

    # ---- inject ----
    html=TPL.replace('__NETDATA__',json.dumps(NET,separators=(',',':'))).replace('__LOGO__','data:image/png;base64,'+LOGO)
    html=html.replace('ZONE 1B · BLOCK 5 · GRID 44N', _pill(zone,blk))
    html=html.replace('Z1B-B5 (44N)', f'{code} (44N)')
    html=html.replace('Initialising Digital Twin \u00b7 Z1B-B5', f'Initialising Digital Twin \u00b7 {code}')
    fn=os.path.join(OUT,f'AmaravatiTwin_{code}.html')
    open(fn,'w').write(html)
    npipe=sum(len(networks[k]) for k in ORDER)
    tot_len=sum(stats[k]['length_m'] for k in ORDER)
    print(f'  {code}: {npipe} pipes, {len(nodes)} manholes, road {nx}x{nz} → {os.path.basename(fn)} ({os.path.getsize(fn)//1024} KB)')
    return {'code':code,'zone':zone,'block':blk,'file':'blocks/'+os.path.basename(fn),
            'pipes':npipe,'manholes':len(nodes),'length_km':round(tot_len/1000,2),
            'lat':round(lat0,6),'lon':round(lon0,6),
            'utils':{k:{'pipes':stats[k]['count'],'km':round(stats[k]['length_m']/1000,2)} for k in ORDER}}

def _zkey(z):
    m=re.match(r'(\d+)(.*)',str(z)); return (int(m.group(1)),m.group(2)) if m else (999,str(z))
def _bkey(b):
    m=re.search(r'(\d+)',str(b)); return int(m.group(1)) if m else 0

files=sorted(glob.glob(f'{UP}/178*.xlsx'))
print('Building',len(files),'blocks...')
man=[]
for p in files:
    try:
        r=build_block(p)
        if r: man.append(r)
    except Exception as e:
        import traceback; print('  ERROR',os.path.basename(p),e); traceback.print_exc()
man.sort(key=lambda m:(_zkey(m['zone']),_bkey(m['block'])))
json.dump(man, open(PKG+'/manifest.json','w'), indent=1)
print('Done:',len(man),'blocks · manifest written')
